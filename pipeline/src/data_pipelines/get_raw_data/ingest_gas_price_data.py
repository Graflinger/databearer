"""Ingest the monthly European gas benchmark and the USD/EUR reference rate.

Run from ``pipeline/``::

    PYTHONPATH=. python src/data_pipelines/get_raw_data/ingest_gas_price_data.py
    PYTHONPATH=. python src/data_pipelines/get_raw_data/ingest_gas_price_data.py --from-raw

Two bounded requests: the World Bank Pink Sheet workbook (CC BY 4.0) and one ECB
SDMX CSV for 2019-01..2025-12. Raw bytes plus a small provenance sidecar are kept
in ignored ``.data/raw/gas_prices/``; ``--from-raw`` reloads them without any
network access. Values are written unchanged to two DuckDB staging tables. Unit
conversion to EUR/MWh happens in dbt (``fact_gas_price_europe_monthly``).
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path

import duckdb
import requests
import yaml
from openpyxl import load_workbook

METADATA = Path("src/config/datasource_metadata/gas_price_sources.yaml")
DEFAULT_RAW_DIR = Path(".data/raw/gas_prices")
DEFAULT_DATABASE = Path(".data/duckdb.db")
PINK_SHEET_FILE = "world_bank_pink_sheet_monthly.xlsx"
ECB_FILE = "ecb_exr_usd_eur_monthly.csv"
TIMEOUT_SECONDS = 30
PERIOD_PINK = re.compile(r"^(\d{4})M(0[1-9]|1[0-2])$")
PERIOD_ECB = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
# The Pink Sheet marks unavailable observations with ellipses.
MISSING_MARKERS = {None, "", "…", "...", ".."}


class SourceError(ValueError):
    """Raised when a source file does not match its documented contract."""


def require(condition, message):
    if not condition:
        raise SourceError(message)


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def load_sources(path: Path = METADATA) -> dict:
    with open(path, encoding="utf-8") as file:
        return yaml.safe_load(file)["sources"]


def download(url: str, max_bytes: int, session=requests) -> bytes:
    """Download one file with a timeout and a hard size cap."""
    with session.get(url, timeout=TIMEOUT_SECONDS, stream=True) as response:
        response.raise_for_status()
        chunks, size = [], 0
        for chunk in response.iter_content(65536):
            size += len(chunk)
            require(size <= max_bytes, f"Response exceeds {max_bytes} bytes: {url}")
            chunks.append(chunk)
    data = b"".join(chunks)
    require(data, f"Empty response: {url}")
    return data


def _number(value, label):
    if isinstance(value, str):
        value = value.strip()
    if value in MISSING_MARKERS:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise SourceError(f"{label}: not a number: {value!r}") from error
    require(math.isfinite(number) and number > 0, f"{label}: expected a positive finite value, got {number}")
    return number


def parse_pink_sheet(raw: bytes, sheet: str, series: str, unit: str) -> dict:
    """Return the Pink Sheet release label and all monthly rows of one series."""
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises several unrelated types
        raise SourceError(f"Pink Sheet is not a readable workbook: {error}") from error
    require(sheet in workbook.sheetnames, f"Pink Sheet has no sheet {sheet!r}")
    rows = list(workbook[sheet].iter_rows(values_only=True))
    workbook.close()

    headers = [index for index, row in enumerate(rows) if series in row]
    require(len(headers) == 1, f"Expected exactly one header row containing {series!r}")
    header = headers[0]
    column = rows[header].index(series)
    require(rows[header].count(series) == 1, f"Series {series!r} appears more than once")
    require(header + 1 < len(rows) and rows[header + 1][column] == unit, f"Unit of {series!r} is not {unit!r}")
    updated = [cell for row in rows[:header] for cell in row if isinstance(cell, str) and cell.startswith("Updated on ")]
    require(len(updated) == 1, "Pink Sheet release label ('Updated on ...') not found")
    release = datetime.strptime(updated[0].removeprefix("Updated on ").strip(), "%B %d, %Y").date().isoformat()

    observations, seen = [], set()
    for row in rows[header + 2:]:
        label = row[0]
        if label is None:
            continue
        match = PERIOD_PINK.match(str(label).strip())
        require(match is not None, f"Unexpected period label {label!r}")
        period = f"{match.group(1)}-{match.group(2)}"
        require(period not in seen, f"Duplicate period {period}")
        seen.add(period)
        observations.append((period, _number(row[column] if column < len(row) else None, f"Pink Sheet {period}")))
    require(observations, "Pink Sheet contains no monthly observations")
    require([period for period, _ in observations] == sorted(seen), "Pink Sheet periods are not ascending")
    return {"release_date": release, "rows": observations}


def parse_ecb_csv(raw: bytes, series_key: str) -> list[tuple[str, float]]:
    """Return monthly USD per EUR observations from an ECB SDMX CSV export."""
    reader = csv.DictReader(io.StringIO(raw.decode("utf-8")))
    require(reader.fieldnames and {"KEY", "TIME_PERIOD", "OBS_VALUE"} <= set(reader.fieldnames), "ECB CSV columns missing")
    rows, seen = [], set()
    for record in reader:
        require(record["KEY"] == series_key, f"Unexpected ECB series {record['KEY']!r}")
        period = record["TIME_PERIOD"].strip()
        require(PERIOD_ECB.match(period) is not None, f"Unexpected ECB period {period!r}")
        require(period not in seen, f"Duplicate ECB period {period}")
        seen.add(period)
        value = _number(record["OBS_VALUE"], f"ECB {period}")
        require(value is not None, f"ECB {period}: missing value")
        rows.append((period, value))
    require(rows, "ECB CSV contains no observations")
    return sorted(rows)


def store_raw(raw_dir: Path, filename: str, data: bytes, url: str) -> dict:
    raw_dir.mkdir(parents=True, exist_ok=True)
    provenance = {
        "filename": filename, "url": url, "sha256": sha256(data), "bytes": len(data),
        "retrieved_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
    }
    temporary = raw_dir / f".{filename}.tmp"
    temporary.write_bytes(data)
    temporary.replace(raw_dir / filename)
    (raw_dir / f"{filename}.json").write_text(json.dumps(provenance, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return provenance


def read_raw(raw_dir: Path, filename: str) -> tuple[bytes, dict]:
    data = (raw_dir / filename).read_bytes()
    provenance = json.loads((raw_dir / f"{filename}.json").read_text(encoding="utf-8"))
    require(provenance["sha256"] == sha256(data) and provenance["bytes"] == len(data), f"Raw file {filename} does not match its sidecar")
    return data, provenance


def load_staging(con, gas: dict, gas_provenance: dict, fx: list, fx_provenance: dict, sources: dict) -> None:
    pink, ecb = sources["world_bank_pink_sheet"], sources["ecb_usd_eur"]
    con.execute("CREATE SCHEMA IF NOT EXISTS staging")
    con.execute("""
        CREATE OR REPLACE TABLE staging.world_bank_gas_europe_monthly (
            period VARCHAR, usd_per_mmbtu DOUBLE, series VARCHAR, unit VARCHAR,
            release_date VARCHAR, source_url VARCHAR, source_sha256 VARCHAR, retrieved_at VARCHAR)""")
    con.executemany(
        "INSERT INTO staging.world_bank_gas_europe_monthly VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        [(period, value, pink["series"], pink["unit"], gas["release_date"], gas_provenance["url"],
          gas_provenance["sha256"], gas_provenance["retrieved_at"]) for period, value in gas["rows"]])
    con.execute("""
        CREATE OR REPLACE TABLE staging.ecb_usd_eur_monthly (
            period VARCHAR, usd_per_eur DOUBLE, series_key VARCHAR,
            source_url VARCHAR, source_sha256 VARCHAR, retrieved_at VARCHAR)""")
    con.executemany(
        "INSERT INTO staging.ecb_usd_eur_monthly VALUES (?, ?, ?, ?, ?, ?)",
        [(period, value, ecb["series_key"], fx_provenance["url"], fx_provenance["sha256"], fx_provenance["retrieved_at"])
         for period, value in fx])


def ingest(database: Path = DEFAULT_DATABASE, raw_dir: Path = DEFAULT_RAW_DIR, from_raw: bool = False, session=requests) -> dict:
    sources = load_sources()
    pink, ecb = sources["world_bank_pink_sheet"], sources["ecb_usd_eur"]
    if from_raw:
        pink_raw, pink_provenance = read_raw(raw_dir, PINK_SHEET_FILE)
        ecb_raw, ecb_provenance = read_raw(raw_dir, ECB_FILE)
    else:
        pink_raw = download(pink["url"], pink["max_bytes"], session)
        ecb_raw = download(ecb["url"], ecb["max_bytes"], session)
    # Parse before storing so that a malformed download never replaces good raw files.
    gas = parse_pink_sheet(pink_raw, pink["sheet"], pink["series"], pink["unit"])
    fx = parse_ecb_csv(ecb_raw, ecb["series_key"])
    if not from_raw:
        pink_provenance = store_raw(raw_dir, PINK_SHEET_FILE, pink_raw, pink["url"])
        ecb_provenance = store_raw(raw_dir, ECB_FILE, ecb_raw, ecb["url"])
    with duckdb.connect(str(database)) as con:
        con.execute("BEGIN")
        load_staging(con, gas, pink_provenance, fx, ecb_provenance, sources)
        con.execute("COMMIT")
    return {"gas_rows": len(gas["rows"]), "fx_rows": len(fx), "release_date": gas["release_date"]}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    parser.add_argument("--raw-dir", type=Path, default=DEFAULT_RAW_DIR)
    parser.add_argument("--from-raw", action="store_true", help="reload stored raw files without network access")
    args = parser.parse_args(argv)
    result = ingest(args.database, args.raw_dir, args.from_raw)
    print(json.dumps(result, sort_keys=True))


if __name__ == "__main__":
    main()
